#!/usr/bin/env python3
"""Teste les cibles exactes en HTTPS et HTTP, sans découverte ni fuzzing.
TXT : une cible par ligne. Dépendance : pip install httpx
Sans option : HTTP seul. --was : HTTP + WAS. --was-only : WAS seul.
Clés WAS : TENABLE_ACCESS_KEY et TENABLE_SECRET_KEY dans l’environnement.
WAS 1 = exécution observée pour l’URL exacte ; 0 = aucune exécution visible.
Cela ne garantit ni un scan terminé ni une couverture complète de l’application.
Identité : protocole, hôte, port effectif, chemin et paramètres exacts.
Seule la racine vide équivaut à / ; /get n’équivaut jamais à /.
TLS non validé pour les cibles ; TLS validé pour l’API Tenable.
Redirections suivies uniquement sur le même hôte (dix maximum).
"""
import argparse
import csv
import io
import ipaddress
import os
import time
import re
import sys

import httpx
from pathlib import Path
from urllib.parse import urlsplit, urlunsplit, urljoin, quote


def load_targets(source, column):
    p = Path(source)
    if not p.is_file():
        if p.suffix.lower() in ('.xlsx', '.xls', '.csv', '.txt') and '://' not in source:
            raise ValueError('Fichier introuvable : ' + source)
        return [source], None
    if p.suffix.lower() == '.xlsx':
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise ValueError('Lecture Excel : installer openpyxl ou exporter en CSV.') from None
        wb = load_workbook(p, read_only=True, data_only=True)
        try:
            values = [r[column] for r in wb.active.iter_rows(values_only=True) if len(r) > column]
        finally:
            wb.close()
    elif p.suffix.lower() == '.xls':
        raise ValueError('Convertir le fichier .xls en .xlsx ou CSV.')
    elif p.suffix.lower() == '.csv':
        text = p.read_text(encoding='utf-8-sig')
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        values = [r[column] for r in csv.reader(io.StringIO(text), dialect) if len(r) > column]
    else:
        values = p.read_text(encoding='utf-8-sig').splitlines()
    return values, p.stem


def prepare(values, errors):
    rows, seen = [], set()
    headers = {'url', 'urls', 'cible', 'cibles', 'domaine', 'domaines', 'target', 'targets'}
    for number, value in enumerate(values, 1):
        s = str(value or '').strip().lstrip('\ufeff')
        if not s or s.lower() in headers or s.startswith('#'):
            continue
        try:
            if any(c.isspace() or ord(c) < 32 or ord(c) == 127 for c in s):
                raise ValueError('espaces/caractères de contrôle')
            try:
                ip = ipaddress.ip_address(s)
                if ip.version == 6:
                    s = '[' + ip.compressed + ']'
            except ValueError:
                pass
            u = urlsplit(s if '://' in s else '//' + s)
            if not u.hostname or u.username is not None or u.password is not None:
                raise ValueError('hôte absent/identifiants dans URL')
            if u.scheme and u.scheme.lower() not in ('http', 'https'):
                raise ValueError('protocole invalide')
            try:
                host = ipaddress.ip_address(u.hostname).compressed
            except ValueError:
                host = u.hostname.encode('idna').decode().lower()
            if any(c in host for c in '/\\?#%') or host.startswith('.') or '..' in host:
                raise ValueError('hôte invalide')
            host = '[' + host + ']' if ':' in host else host
            port = u.port
            if port is not None and not 1 <= port <= 65535:
                raise ValueError('port invalide')
            if not (u.path + u.query).isascii():
                raise ValueError('encoder chemin/paramètres non ASCII en %XX')
            identity = (host, port, u.path or '/', u.query)
            if identity in seen:
                continue
            seen.add(identity)
            label = host + u.path + ('?' + u.query if u.query else '')
            authority = host + (':' + str(port) if port else '')
            urls = [urlunsplit((scheme, authority, u.path or '/', u.query, '')) for scheme in ('https', 'http')]
            rows.append((label, [port or 443, port or 80], urls))
        except (ValueError, UnicodeError) as exc:
            errors.append(f'Entrée {number} ignorée : {exc}')
    if not rows:
        raise ValueError('Aucune cible valide.')
    return rows


def response(client, url, errors):
    codes, current, seen = [], url, set()
    original_host = urlsplit(url).hostname
    client.cookies.clear()  # Chaque protocole repart sans session précédente.
    try:
        for hop in range(11):
            normalized = str(httpx.URL(current))
            if normalized in seen:
                errors.append(url + ' : boucle de redirection')
                return '/'.join(codes + ['BOUCLE_REDIRECTION'])
            seen.add(normalized)
            # Seuls les en-têtes sont nécessaires : ne pas télécharger le corps.
            with client.stream('GET', current) as result:
                codes.append(str(result.status_code))
                location = result.headers.get('location')
            if result.status_code not in (301, 302, 303, 307, 308) or not location:
                return '/'.join(codes)
            destination = urljoin(current, location)
            parts = urlsplit(destination)
            if (parts.scheme not in ('http', 'https') or parts.hostname != original_host
                    or parts.username is not None or parts.password is not None):
                errors.append(url + ' : redirection hors de l’hôte ciblé ou destination invalide : ' + destination)
                return '/'.join(codes + ['REDIRECTION_HORS_PERIMETRE'])
            if hop == 10:
                errors.append(url + ' : limite de dix redirections atteinte')
                return '/'.join(codes + ['LIMITE_REDIRECTIONS'])
            current = destination
    except (httpx.HTTPError, httpx.InvalidURL, ValueError) as exc:
        message = str(exc)
        errors.append(url + ' : ' + type(exc).__name__ + ': ' + message)
        low = message.lower()
        if isinstance(exc, httpx.TimeoutException):
            reason = 'TIMEOUT'
        elif any(x in low for x in ('tls', 'ssl', 'x509', 'handshake')):
            reason = 'ERREUR_TLS'
        elif 'refused' in low:
            reason = 'CONNEXION_REFUSEE'
        else:
            reason = 'ERREUR'
        return '/'.join(codes + [reason])


def was_request(client, path, params):
    for attempt in range(3):
        result = client.post(path, params=params, json={})
        if result.status_code not in (429, 502, 503, 504) or attempt == 2:
            break
        wait = result.headers.get('Retry-After', '2')
        time.sleep(min(int(wait) if wait.isdigit() else 2, 10))
    if result.status_code != 200:
        raise ValueError(f'API Tenable HTTP {result.status_code} : vérifier les clés, les droits et le service.')
    data = result.json()
    if not isinstance(data, dict) or not isinstance(data.get('items'), list):
        raise ValueError('Réponse API Tenable inattendue : champ items absent/non valide.')
    return data


def execution_state(config):
    if 'last_scan' not in config:
        return 'A_VERIFIER', {}, 'Information de dernière exécution absente.'
    last = config['last_scan']
    if last is None:
        return 'JAMAIS_EXECUTE', {}, 'Aucune dernière exécution enregistrée.'
    if not isinstance(last, dict) or not last:
        return 'A_VERIFIER', {}, 'Dernière exécution non interprétable.'
    status = str(last.get('status') or '').lower()
    metadata = last.get('metadata') or {}
    if not isinstance(metadata, dict):
        return 'A_VERIFIER', last, 'Métadonnées d’exécution non interprétables.'
    traffic = any(isinstance(metadata.get(k), (int, float)) and metadata[k] > 0
                  for k in ('request_count', 'audited_urls', 'audited_pages'))
    if last.get('started_at') or traffic or status in ('running', 'paused', 'stopping', 'completed'):
        return 'DEJA_EXECUTE', last, 'Exécution observée ; ne signifie pas nécessairement un scan terminé avec succès.'
    if status in ('pending', 'queued', 'scheduled', 'initializing'):
        return 'EN_ATTENTE', last, 'Exécution prévue/en attente, démarrage non confirmé.'
    return 'A_VERIFIER', last, 'Un historique existe, mais le démarrage du scan n’est pas confirmé.'


def url_identity(url):
    """Pas de comparaison par nom, préfixe, casse du chemin ou destination redirigée."""
    u = urlsplit(url)
    if u.scheme.lower() not in ('http', 'https') or not u.hostname or u.username or u.password:
        raise ValueError('Cible WAS non interprétable : URL HTTP/HTTPS complète requise.')
    return (u.scheme.lower(), u.hostname.encode('idna').decode().lower(),
            u.port or (443 if u.scheme.lower() == 'https' else 80), u.path or '/', u.query)


def was_pages(client, path, sort):
    offset, seen = 0, set()
    while True:
        page = was_request(client, path, {'limit': 200, 'offset': offset, 'sort': sort})
        total = page.get('pagination', {}).get('total')
        items = page['items']
        if not items:
            if isinstance(total, int) and offset < total:
                raise ValueError('Pagination WAS incomplète ; aucun statut fiable.')
            break
        for item in items:
            if not isinstance(item, dict):
                raise ValueError('Élément WAS non interprétable.')
            identifier = item.get('scan_id') if '/scans/' in path else item.get('config_id')
            if not identifier or identifier in seen:
                raise ValueError('Identifiant WAS absent ou pagination répétée.')
            seen.add(identifier)
            yield item
        offset += len(items)
        if isinstance(total, int) and offset >= total:
            break


def check_tenable(rows, base_url):
    access, secret = os.environ.get('TENABLE_ACCESS_KEY'), os.environ.get('TENABLE_SECRET_KEY')
    if not access or not secret:
        raise ValueError('Définir TENABLE_ACCESS_KEY et TENABLE_SECRET_KEY.')
    base = urlsplit(base_url)
    if base.scheme != 'https' or not base.hostname or base.username or base.password or base.query or base.fragment or base.path not in ('', '/'):
        raise ValueError('--tenable-url doit être une origine HTTPS sans chemin.')
    wanted = {url_identity(url) for _, _, urls in rows for url in urls}
    executed = set()
    headers = {'X-ApiKeys': f'accessKey={access};secretKey={secret}', 'Accept': 'application/json'}
    with httpx.Client(base_url=base_url.rstrip('/'), headers=headers, verify=True,
                      timeout=30, trust_env=False, follow_redirects=False) as client:
        configs = list(was_pages(client, '/was/v2/configs/search', 'name:asc'))
        # Historique complet : une configuration peut avoir changé de cible,
        # ou sa dernière exécution être en attente malgré un ancien scan exécuté.
        for config in configs:
            if config.get('last_scan', 'missing') is None:
                continue
            path = '/was/v2/configs/' + quote(str(config['config_id']), safe='') + '/scans/search'
            history_found = False
            for scan in was_pages(client, path, 'created_at:desc'):
                history_found = True
                state, _, _ = execution_state({'last_scan': scan})
                target = scan.get('target')
                if state == 'A_VERIFIER':
                    raise ValueError('Historique WAS ambigu ; aucun classement 0/1 fiable.')
                if state != 'DEJA_EXECUTE':
                    continue
                if not isinstance(target, str) or not target:
                    raise ValueError('Cible historique WAS absente ; comparaison exacte impossible.')
                identity = url_identity(target)
                if identity in wanted:
                    executed.add(identity)
            if not history_found and config.get('last_scan'):
                raise ValueError('Historique WAS absent malgré une dernière exécution ; relancer.')
    return {url: int(url_identity(url) in executed) for _, _, urls in rows for url in urls}


def write_rows(writer, rows, client, states, errors, was_only=False):
    for domain, ports, urls in rows:
        codes = ['',''] if was_only else [response(client, url, errors) for url in urls]
        status = [states[url] for url in urls] if states is not None else [None, None]
        if codes[0] == codes[1] and status[0] == status[1]:
            values = [domain] if was_only else [domain, '/'.join(dict.fromkeys(map(str, ports))), codes[0]]
            writer.writerow(([status[0]] if states is not None else []) + values)
        else:
            for url, port, code, state in zip(urls, ports, codes, status):
                # Sans colonne port (WAS seul), conserver le protocole pour lever
                # l’ambiguïté. Même précaution si un port explicite est partagé.
                label = url if was_only or ports[0] == ports[1] else domain
                values = [label] if was_only else [label, port, code]
                writer.writerow(([state] if states is not None else []) + values)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('source', help='URL, fichier XLSX, CSV ou TXT')
    parser.add_argument('column', nargs='?', type=int, default=1, help='colonne Excel/CSV (défaut : 1)')
    modes = parser.add_mutually_exclusive_group()
    modes.add_argument('--was', '--tenable', dest='was', action='store_true', help='HTTP + WAS exact')
    modes.add_argument('--was-only', '--tenable-only', dest='was_only', action='store_true', help='WAS seul, sans requêtes aux cibles')
    parser.add_argument('--tenable-url', default='https://cloud.tenable.com', help='origine de l’API Tenable cloud')
    args = parser.parse_args()
    if args.column < 1:
        parser.error('Colonne minimale : 1')
    errors = []
    try:
        values, name = load_targets(args.source, args.column - 1)
        rows = prepare(values, errors)
        if name is None:
            name = rows[0][0].split('/')[0].split('?')[0].strip('[]')
        name = re.sub(r'[^A-Za-z0-9._-]', '_', name).strip('.') or 'target'
        with_was = args.was or args.was_only
        states = check_tenable(rows, args.tenable_url) if with_was else None
        suffix = '_was_status.csv' if args.was_only else '_http_was_status.csv' if with_was else '_result.csv'
        destination, log_path = Path(name + suffix), Path(name + '_error.log')
        # Écriture temporaire : ne pas laisser un CSV partiel après un échec.
        temporary = destination.with_suffix('.csv.tmp')
        try:
            with temporary.open('w', encoding='utf-8-sig', newline='') as output:
                writer = csv.writer(output)
                fields = ['domaine'] if args.was_only else ['domaine', 'port', 'reponse_http']
                writer.writerow((['was status'] if with_was else []) + fields)
                if args.was_only:
                    write_rows(writer, rows, None, states, errors, was_only=True)
                else:
                    with httpx.Client(verify=False, timeout=5.0, follow_redirects=False, trust_env=False) as client:
                        write_rows(writer, rows, client, states, errors)
            temporary.replace(destination)
        finally:
            temporary.unlink(missing_ok=True)
        log_path.write_text('\n'.join(errors) + ('\n' if errors else ''), encoding='utf-8')
        print(destination.resolve())
        return 0
    except (OSError, ValueError, httpx.HTTPError) as exc:
        print('Erreur : ' + str(exc), file=sys.stderr)
        return 1


if __name__ == '__main__':
    raise SystemExit(main())
